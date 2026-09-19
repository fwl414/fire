"""统计报表导出服务

统一使用 report_statistics_service 的真实租户统计口径生成 CSV/XLSX，
生成后由 report_archive_service 落盘并登记，供受控下载与追溯。
"""
from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy.orm import Session

from database import ReportArtifact
from services.report_archive_service import archive_report
from services.report_statistics_service import VALID_REPORT_TYPES, build_report_statistics

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# 对外格式名 -> 内部格式名
FORMAT_ALIASES = {
    "csv": "csv",
    "xlsx": "xlsx",
    "excel": "xlsx",
    "xls": "xlsx",
}

TITLE_PREFIX = "消防"


@dataclass
class ExportResult:
    artifact: ReportArtifact
    content: bytes
    media_type: str
    filename: str


def _trend_columns(trend: List[Dict[str, Any]]) -> Tuple[List[str], List[List[Any]]]:
    """趋势数据结构在不同报表类型下不同（按日期或按维度），统一成表头+行。"""
    if not trend:
        return [], []
    if isinstance(trend[0], dict) and "date" in trend[0]:
        return ["日期", "数量"], [[r.get("date", ""), r.get("count", 0)] for r in trend]
    return ["名称", "数值"], [[r.get("name", ""), r.get("value", 0)] for r in trend]


def _report_title(report_type: str) -> str:
    names = {
        "inspection": "巡检统计报表",
        "alert": "告警统计报表",
        "workorder": "工单统计报表",
        "device": "设备状态报表",
        "risk": "风险分析报表",
    }
    return f"{TITLE_PREFIX}{names.get(report_type, '统计报表')}"


def render_csv(stats: Dict[str, Any], report_type: str, title: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([title])
    writer.writerow(["报表类型", report_type])
    writer.writerow(["统计周期", f"{stats.get('start_date', '')} ~ {stats.get('end_date', '')}"])
    writer.writerow(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    writer.writerow([])

    writer.writerow(["汇总数据"])
    for key, value in (stats.get("summary") or {}).items():
        writer.writerow([key, value])
    writer.writerow([])

    headers, rows = _trend_columns(stats.get("trend") or [])
    if rows:
        writer.writerow(["趋势统计"])
        writer.writerow(headers)
        writer.writerows(rows)
        writer.writerow([])

    by_category = stats.get("byCategory") or []
    if by_category:
        writer.writerow(["分类统计"])
        writer.writerow(["名称", "数量"])
        for item in by_category:
            writer.writerow([item.get("name", ""), item.get("value", 0)])

    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(stats: Dict[str, Any], report_type: str, title: str) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:  # pragma: no cover - 依赖缺失时的明确报错
        raise RuntimeError("缺少 openpyxl 依赖，无法导出 Excel") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:D1")
    ws["A2"] = f"报表类型：{report_type}"
    ws["A3"] = f"统计周期：{stats.get('start_date', '')} ~ {stats.get('end_date', '')}"
    ws["A4"] = f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    row = 6
    ws.cell(row=row, column=1, value="汇总数据").font = Font(bold=True, size=12)
    row += 1
    for key, value in (stats.get("summary") or {}).items():
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=value)
        row += 1

    headers, rows = _trend_columns(stats.get("trend") or [])
    if rows:
        ws2 = wb.create_sheet("趋势统计")
        for col, header in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        for r, data_row in enumerate(rows, 2):
            for col, value in enumerate(data_row, 1):
                ws2.cell(row=r, column=col, value=value)

    by_category = stats.get("byCategory") or []
    if by_category:
        ws3 = wb.create_sheet("分类统计")
        ws3.cell(row=1, column=1, value="名称").font = Font(bold=True)
        ws3.cell(row=1, column=2, value="数量").font = Font(bold=True)
        for r, item in enumerate(by_category, 2):
            ws3.cell(row=r, column=1, value=item.get("name", ""))
            ws3.cell(row=r, column=2, value=item.get("value", 0))

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_report(
    db: Session,
    *,
    tenant_id: Optional[int],
    payload: Dict[str, Any],
    user_id: Optional[int] = None,
    user_name: str = "",
) -> ExportResult:
    """生成报表：真实租户数据 → 渲染 → 落盘登记。"""
    report_type = str(payload.get("report_type", "inspection") or "inspection")
    if report_type not in VALID_REPORT_TYPES:
        raise ValueError(f"不支持的报表类型：{report_type}，可选值：{', '.join(VALID_REPORT_TYPES)}")

    raw_format = str(payload.get("format", "xlsx") or "xlsx").lower()
    fmt = FORMAT_ALIASES.get(raw_format)
    if not fmt:
        raise ValueError(f"不支持的导出格式：{raw_format}，可选值：{', '.join(sorted(FORMAT_ALIASES))}")

    period = str(payload.get("period", "month") or "month")
    stats = build_report_statistics(
        db,
        tenant_id,
        report_type=report_type,
        period=period,
        start_date=payload.get("start_date", "") or "",
        end_date=payload.get("end_date", "") or "",
    )

    title = _report_title(report_type)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{TITLE_PREFIX}_{report_type}_report_{timestamp}.{fmt}"

    content = render_csv(stats, report_type, title) if fmt == "csv" else render_xlsx(stats, report_type, title)
    media_type = MEDIA_TYPES[fmt]

    artifact = archive_report(
        db,
        tenant_id=tenant_id,
        report_type=report_type,
        report_format=fmt,
        title=title,
        filename=filename,
        content=content,
        media_type=media_type,
        params={
            "period": period,
            "start_date": stats.get("start_date", ""),
            "end_date": stats.get("end_date", ""),
            "requested_format": raw_format,
        },
        generated_by=user_id,
        generated_by_name=user_name,
    )

    return ExportResult(
        artifact=artifact,
        content=content,
        media_type=media_type,
        filename=filename,
    )
